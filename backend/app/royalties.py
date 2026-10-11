"""Parsing of store sales/royalty reports (CSV or XLSX) into normalized rows.

Stores export different layouts, so columns are detected by header synonyms (EN/PL); when detection fails
the caller gets the header list and can send an explicit mapping.
"""
import csv
import io
import re
import zipfile
from datetime import date, datetime

FIELDS = {
    "title": ["title", "product title", "book title", "tytuł", "tytul", "nazwa produktu", "produkt"],
    "id": ["asin", "isbn", "asin/isbn", "eisbn", "primary isbn", "identifier", "ean", "product id", "identyfikator"],
    "period": ["royalty date", "date", "transaction date", "earnings date", "sale date", "period", "month", "invoice date",
               "data", "okres", "miesiąc", "data sprzedaży"],
    "units": ["net units sold", "units sold", "units", "quantity", "qty", "sztuk", "ilość", "liczba sztuk"],
    "amount": ["royalty", "earnings amount", "net earnings", "publisher revenue", "your earnings", "net revenue", "amount",
               "earnings", "proceeds", "tantiemy", "kwota", "przychód", "wynagrodzenie", "honorarium"],
    "currency": ["currency", "earnings currency", "payout currency", "royalty currency", "waluta"],
    "market": ["marketplace", "country", "buyer country", "country of sale", "store", "territory", "rynek", "kraj"],
}
REQUIRED = ("title", "amount")
LABELS = {"title": "tytuł", "id": "ISBN/ASIN", "period": "data", "units": "sztuki", "amount": "kwota", "currency": "waluta", "market": "rynek"}
MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}
MARKET_CURRENCY = {"amazon.com": "USD", "amazon.co.uk": "GBP", "amazon.de": "EUR", "amazon.fr": "EUR", "amazon.es": "EUR",
                   "amazon.it": "EUR", "amazon.nl": "EUR", "amazon.pl": "PLN", "amazon.ca": "CAD", "amazon.com.au": "AUD"}
MAX_ROWS = 50_000


class ReportError(ValueError):
    def __init__(self, message: str, headers: list[str] | None = None):
        super().__init__(message)
        self.headers = headers or []


def _norm(h) -> str:
    return re.sub(r"\s+", " ", str(h or "").strip().lower().replace("_", " "))


def read_table(data: bytes, filename: str) -> list[list]:
    """Rows (header row included) from CSV (any common delimiter/encoding) or XLSX (first sheet with data)."""
    if data[:2] == b"PK" or filename.lower().endswith(".xlsx"):
        from openpyxl import load_workbook  # defusedxml installed -> hardened XML parsing

        with zipfile.ZipFile(io.BytesIO(data)) as z:
            if sum(i.file_size for i in z.infolist()) > 200 * 1024 * 1024:
                raise ReportError("Plik jest zbyt duży po rozpakowaniu")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 50)]
            if sum(1 for r in rows if any(c not in (None, "") for c in r)) > 1:
                return rows
        raise ReportError("Arkusz nie zawiera danych")
    for enc in ("utf-8-sig", "cp1250", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [r for r in csv.reader(io.StringIO(text), dialect)][: MAX_ROWS + 50]


def find_header(rows: list[list]) -> int:
    """Reports often start with a few title lines: the header is the first row naming a title and an amount column."""
    for i, r in enumerate(rows[:30]):
        names = {_norm(c) for c in r}
        if any(n in names for n in FIELDS["title"]) and any(n in names for n in FIELDS["amount"]):
            return i
    return 0


def detect(headers: list[str]) -> dict[str, int]:
    norm = [_norm(h) for h in headers]
    found = {}
    for field, names in FIELDS.items():
        for name in names:  # synonyms are ordered by preference
            if name in norm:
                found[field] = norm.index(name)
                break
    return found


def parse_number(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d,.\-]", "", str(v or ""))
    if not s or s in "-.,":
        return 0.0
    if "," in s and "." in s:  # the later separator is the decimal one
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".") if len(s.split(",")[-1]) != 3 else s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return 0.0


def parse_period(v) -> str | None:
    if isinstance(v, (datetime, date)):
        return f"{v.year:04d}-{v.month:02d}"
    s = str(v or "").strip().lower()
    if m := re.search(r"(\d{4})[-/.](\d{1,2})", s):
        y, mo = int(m.group(1)), int(m.group(2))
    elif m := re.search(r"(\d{1,2})[-/.](\d{1,2})[-/.](\d{4})", s):
        a, b, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        mo = b if a > 12 or b <= 12 else a  # day-first unless that is impossible (US order)
    elif m := re.search(r"([a-z]{3})[a-z]*[ ,-]+(\d{4})", s):
        y, mo = int(m.group(2)), MONTHS.get(m.group(1), 0)
    else:
        return None
    return f"{y:04d}-{mo:02d}" if 1 <= mo <= 12 and 2000 <= y <= 2100 else None


def normalize(rows: list[list], mapping: dict[str, int] | None, default_period: str | None, default_currency: str) -> list[dict]:
    if not rows:
        raise ReportError("Plik jest pusty")
    h = find_header(rows)
    headers = [str(c or "").strip() for c in rows[h]]
    cols = mapping or detect(headers)
    missing = [LABELS[f] for f in REQUIRED if f not in cols]
    if missing:
        raise ReportError(f"Nie rozpoznano kolumn: {', '.join(missing)}. Wskaż je ręcznie.", headers)
    if "period" not in cols and not default_period:
        raise ReportError("Raport nie ma kolumny z datą — wybierz miesiąc raportu albo wskaż kolumnę daty.", headers)
    out = []
    for r in rows[h + 1:]:
        if not any(c not in (None, "") for c in r):
            continue

        def cell(f):
            i = cols.get(f)
            return r[i] if i is not None and i < len(r) else None
        title = str(cell("title") or "").strip()
        if not title or _norm(title) in ("total", "suma", "razem", "totals"):
            continue
        period = parse_period(cell("period")) if "period" in cols else None
        market = str(cell("market") or "").strip()[:40]
        currency = str(cell("currency") or "").strip().upper()[:3] or MARKET_CURRENCY.get(market.lower(), default_currency)
        out.append({"title": title[:500], "id": re.sub(r"[^0-9A-Za-z]", "", str(cell("id") or ""))[:40],
                    "period": period or default_period, "units": int(round(parse_number(cell("units")))),
                    "amount": int(round(parse_number(cell("amount")) * 100)), "currency": currency, "market": market})
    if not out:
        raise ReportError("Nie znaleziono wierszy ze sprzedażą", headers)
    if any(not x["period"] for x in out):
        raise ReportError("Część wierszy nie ma czytelnej daty — wybierz miesiąc raportu.", headers)
    return out
